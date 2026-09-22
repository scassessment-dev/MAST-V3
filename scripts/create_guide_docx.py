import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DATA_PATH = ROOT / "tmp" / "pdfs" / "guide-data.json"
DOCX_PATH = ROOT / "tmp" / "pdfs" / "MAST_Guide_Gujarati_Document.docx"

FONT = "Shruti"
FONT_BOLD = "Shruti"
SECTIONS = [
    ("પરિચય", "introduction"),
    ("કેવી રીતે ઓળખશો", "recognition"),
    ("શક્તિઓ", "strengths"),
    ("ધ્યાન રાખવું", "cautions"),
    ("પૂ. સંત માટે", "saintGuidance"),
    ("વિકાસ અને સેવા", "growth"),
    ("દબાણ આવે ત્યારે", "pressure"),
]


def set_font(run, size=None, bold=None, color=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run._element.rPr.rFonts.set(qn("w:cs"), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    keep = OxmlElement("w:keepNext")
    p_pr.append(keep)


def add_text(document, text, *, bold=False, style=None, size=12.5, space_after=7, indent=0):
    paragraph = document.add_paragraph(style=style)
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = 1.35
    paragraph.paragraph_format.left_indent = Cm(indent)
    run = paragraph.add_run(text)
    set_font(run, size=size, bold=bold, color=(0, 0, 0))
    return paragraph


def add_heading(document, text, level):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(13 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(5)
    set_keep_with_next(paragraph)
    run = paragraph.add_run(text)
    set_font(run, size=19 if level == 1 else 15, bold=True, color=(0, 0, 0))
    return paragraph


def service_label(service):
    short_code = service["shortCode"].strip()
    name = service["name"].strip()
    spelled_out_name = any(ch.isalpha() and ch.isascii() for ch in short_code) and "-" in short_code
    if short_code == name or name.startswith(short_code + " ") or spelled_out_name:
        return name
    return f"{short_code} - {name}"


def add_service_list(document, services, mast_type, title):
    add_heading(document, f"{mast_type} - {title} માટે સંભવિત સેવાઓ", 2)
    add_text(document, "આ માત્ર માર્ગદર્શક મેળ છે. વ્યક્તિની આવડત, તૈયારી અને પૂ. સંતનો નિર્ણય વધુ મહત્વનો છે.", size=11.5, space_after=7)
    grouped = {}
    for service in services:
        grouped.setdefault(service["channel"], []).append(service)
    for channel, items in grouped.items():
        paragraph = add_text(document, channel, bold=True, size=12, space_after=2)
        set_keep_with_next(paragraph)
        add_text(document, " • ".join(service_label(item) for item in items), size=11.5, space_after=6)


def main():
    data = json.loads(DATA_PATH.read_text(encoding="utf-8"))
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.9)
    section.left_margin = Cm(2.25)
    section.right_margin = Cm(2.25)

    normal = document.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal._element.rPr.rFonts.set(qn("w:cs"), FONT)
    normal.font.size = Pt(12.5)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run("MAST માર્ગદર્શિકા")
    set_font(run, size=25, bold=True, color=(0, 0, 0))

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(20)
    subtitle.paragraph_format.line_spacing = 1.35
    run = subtitle.add_run("પ્રકૃતિને સમજીએ, વ્યક્તિને બાંધીએ નહીં")
    set_font(run, size=14, bold=False, color=(0, 0, 0))

    add_text(document, "આ દસ્તાવેજમાં M, A, S અને T ચારેય MAST પ્રકૃતિનું સંપૂર્ણ માર્ગદર્શન, તેમના મુખ્ય શબ્દો, quotes, ઓળખ, શક્તિઓ, ધ્યાન રાખવાની બાબતો, પૂ. સંત માટે માર્ગદર્શન, વિકાસ અને સેવા, દબાણ સમયે વર્તન, સંભવિત સેવાઓ તથા જરૂરી વિવેકના મુદ્દાઓ સામેલ છે.", size=12.5, space_after=13)

    add_heading(document, "વિષયસૂચિ", 1)
    for key in ["M", "A", "S", "T"]:
        guide = data["guides"][key]
        add_text(document, f"{key} - {guide['title']} : {guide['subtitle']}", size=12, space_after=2, indent=0.3)
    add_text(document, "જરૂરી વિવેક", size=12, space_after=12, indent=0.3)

    for index, key in enumerate(["M", "A", "S", "T"]):
        guide = data["guides"][key]
        if index > 0:
            document.add_page_break()
        add_heading(document, f"{key} - {guide['title']}", 1)
        add_text(document, guide["subtitle"], bold=True, size=13, space_after=8)
        quote = document.add_paragraph()
        quote.paragraph_format.left_indent = Cm(0.5)
        quote.paragraph_format.space_after = Pt(7)
        run = quote.add_run(f"“{guide['quote']}”")
        set_font(run, size=13, bold=True, color=(0, 0, 0))
        add_text(document, "મુખ્ય શબ્દો: " + " | ".join(guide["keywords"]), bold=True, size=12, space_after=10)
        for number, (heading, field) in enumerate(SECTIONS, 1):
            add_heading(document, f"{number}. {heading}", 2)
            for text in guide[field]:
                add_text(document, text, size=12.5)
        add_service_list(document, [service for service in data["services"] if service["probablePrimaryType"] == key], key, guide["title"])
        add_heading(document, "યાદ રાખવા જેવું", 2)
        add_text(document, guide["closing"], bold=True, size=12.5, space_after=10)

    document.add_page_break()
    vivek = data["vivek"]
    add_heading(document, vivek["title"], 1)
    add_text(document, vivek["subtitle"], bold=True, size=13, space_after=10)
    add_heading(document, "વાંચતાં પહેલાં", 2)
    for text in vivek["introduction"]:
        add_text(document, text, size=12.5)
    for point in vivek["points"]:
        add_heading(document, f"{point['number']}. {point['title']}", 2)
        for text in point["paragraphs"]:
            add_text(document, text, size=12.5)
    add_heading(document, "છેલ્લે યાદ રાખવું", 2)
    add_text(document, vivek["closing"], bold=True, size=12.5, space_after=10)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(4)
    run = footer.add_run("MAST માર્ગદર્શિકા")
    set_font(run, size=9.5, color=(80, 80, 80))

    DOCX_PATH.parent.mkdir(parents=True, exist_ok=True)
    document.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    main()
